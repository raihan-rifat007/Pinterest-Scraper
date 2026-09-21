from __future__ import annotations

import asyncio
import io
import json
import os
import queue
import re
import threading
import time
import uuid
import zipfile
from pathlib import Path

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.openapi.utils import get_openapi
from pydantic import BaseModel, ConfigDict, Field

from core.dedupe import DedupeStore
from core.downloader import download_all
from core.http import build_session
from core.scraper import (
    board_pins, enrich_with_details, related_pins,
    search_pins, typeahead_suggestions,
)
from core.storage import save_outputs

STATIC_DIR = Path(__file__).parent.parent / "static"


class ScrapeRequest(BaseModel):
    model_config = ConfigDict(extra="ignore")

    mode: str = Field(default="search", pattern="^(search|board)?$")
    query: str = Field(min_length=1)
    limit: int = Field(default=25, ge=1, le=500)
    download: bool = True
    details: bool = True
    dedup: bool = False
    workers: int = Field(default=4, ge=1, le=16)
    delay: float = Field(default=1.0, ge=0, le=30)
    jitter: float = Field(default=0.5, ge=0, le=10)
    batch_size: int = Field(default=10, ge=1, le=100)
    min_width: int = Field(default=0, ge=0, le=10000)
    min_height: int = Field(default=0, ge=0, le=10000)
    proxy: str = ""


class Job:
    def __init__(self, req: ScrapeRequest, out_dir: Path):
        self.id = uuid.uuid4().hex[:12]
        self.req = req
        self.out_dir = out_dir
        self.out_dir.mkdir(parents=True, exist_ok=True)
        self.events: queue.Queue = queue.Queue()
        self.status = "queued"
        self.phase = ""
        self.pins: list[dict] = []
        self.stats: dict = {}
        self.error = ""
        self.cancelled = False
        self.done = threading.Event()

    def emit(self, **ev):
        self.events.put(json.dumps(ev, default=str, ensure_ascii=False))

    def _check_cancel(self):
        if self.cancelled:
            raise RuntimeError("cancelled by user")

    def run(self):
        try:
            self._run()
            self.status = "cancelled" if self.cancelled else "finished"
        except Exception as e:
            self.status = "cancelled" if "cancelled" in str(e) else "error"
            if self.status == "error":
                self.error = str(e)
        self.done.set()
        self.emit(event="done", status=self.status, error=self.error,
                  total=len(self.pins), stats=self.stats)

    def _run(self):
        req = self.req
        session = build_session(
            proxy_pool=[p.strip() for p in req.proxy.split(",") if p.strip()] or None
        )
        store = (
            DedupeStore(self.out_dir / ".seen_pins.json", scan_dir=self.out_dir)
            if req.dedup else None
        )
        queries = [q.strip() for q in req.query.split(",") if q.strip()] or [req.query]
        batch = len(queries) > 1
        stem = (queries[0] if not batch else queries[0] + "-batch")
        stem = stem.strip().replace(" ", "_")[:40] or "pins"

        collect_total = req.limit * len(queries)
        self.emit(event="phase", phase="collect", total=collect_total, message="Collecting pins")

        def batch_save(partial: list[dict]) -> None:
            self._check_cancel()
            new = [p for p in partial if p["pin_id"] not in {q["pin_id"] for q in self.pins}]
            if store:
                new = store.filter(new)
            self.pins.extend(new)
            self.emit(event="progress", phase="collect", count=len(self.pins), total=collect_total)
            save_outputs(self.pins, self.out_dir, stem)

        def run_one(query: str) -> list[dict]:
            self.emit(event="query_start", query=query,
                      index=queries.index(query) + 1, total=len(queries))
            if req.mode == "board":
                return board_pins(session, query, req.limit,
                                  delay=req.delay, save_cb=batch_save,
                                  batch_size=req.batch_size, jitter=req.jitter)
            return search_pins(session, query, req.limit,
                               delay=req.delay, save_cb=batch_save,
                               batch_size=req.batch_size, jitter=req.jitter)

        existing_ids = {p["pin_id"] for p in self.pins}
        raw_pins: list[dict] = []
        for q in queries:
            for p in run_one(q):
                if p["pin_id"] not in existing_ids:
                    existing_ids.add(p["pin_id"])
                    raw_pins.append(p)
            self._check_cancel()

        pins = list(raw_pins)

        if not pins:
            clean_stem = re.sub(r"[^\w\-]+", "_", stem.lower().strip())
            for candidate in (f"{stem}.json", f"{clean_stem}.json", f"{stem.lower()}.json"):
                cached = self.out_dir / candidate
                if cached.exists():
                    try:
                        loaded = json.loads(cached.read_text(encoding="utf-8"))
                        if isinstance(loaded, list) and loaded:
                            pins = loaded
                            break
                    except Exception:
                        pass

        img_dir = self.out_dir / "images"
        if img_dir.exists():
            for p in pins:
                if not p.get("local_file"):
                    for ext in (".jpg", ".png", ".webp", ".jpeg"):
                        candidate = img_dir / f"{p['pin_id']}{ext}"
                        if candidate.is_file() and candidate.stat().st_size > 0:
                            p["local_file"] = candidate.name
                            break

        self.pins = pins

        if not pins:
            self.emit(event="nothing_new", total=0)
            return

        if req.details:
            self.emit(event="phase", phase="details", total=len(pins))
            done = {"n": 0}

            def _detail_cb(n):
                if self.cancelled:
                    raise RuntimeError("cancelled by user")
                done["n"] = max(done["n"], n)
                self.emit(event="progress", phase="details", count=done["n"], total=len(pins))

            enrich_with_details(session, pins, delay=req.delay,
                                workers=req.workers, jitter=req.jitter,
                                progress_cb=_detail_cb)
            self._check_cancel()

        if req.download:
            self.emit(event="phase", phase="download", total=len(pins))

            def _dl_cb(n):
                if self.cancelled:
                    raise RuntimeError("cancelled by user")
                self.emit(event="progress", phase="download", count=n, total=len(pins))

            self.stats = download_all(
                session, pins, self.out_dir / "images",
                req.workers, req.min_width, req.min_height,
                progress_cb=_dl_cb,
            )
            self._check_cancel()

        summary = save_outputs(pins, self.out_dir, stem)
        if store:
            store.add(pins)
            store.save()
        self.pins = pins
        self.emit(event="saved",
                  json_file=str(summary.get("json", "")),
                  csv_file=str(summary.get("csv", "")))


JOBS: dict[str, Job] = {}


def _out_dir() -> Path:
    return Path("web_output")


app = FastAPI(
    title="Pinterest Scraper API",
    version="1.4.0",
    description="High-quality Pinterest scraper API with job events, exports, and gallery management.",
    docs_url=None,
    redoc_url=None,
)


def _custom_openapi():
    if app.openapi_schema:
        return app.openapi_schema
    schema = get_openapi(
        title=app.title,
        version=app.version,
        description=app.description,
        routes=app.routes,
    )
    app.openapi_schema = schema
    return app.openapi_schema


app.openapi = _custom_openapi


app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


DOCS_HTML = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0, viewport-fit=cover">
<meta name="theme-color" content="#E60023">
<meta name="description" content="Pinterest Scraper API reference — endpoints, schemas, and live try-it-out.">
<meta property="og:type" content="website">
<meta property="og:title" content="Pinterest Scraper — API Reference">
<meta property="og:description" content="High-quality Pinterest scraper API with job events, exports, and gallery.">
<meta property="og:image" content="/static/og-image.png">
<meta property="og:image:width" content="1200">
<meta property="og:image:height" content="630">
<meta name="twitter:card" content="summary_large_image">
<meta name="twitter:title" content="Pinterest Scraper — API Reference">
<meta name="twitter:description" content="High-quality Pinterest scraper API.">
<meta name="twitter:image" content="/static/og-image.png">
<title>Pinterest Scraper · API Reference</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800;900&family=JetBrains+Mono:wght@400;500;600;700&display=swap" rel="stylesheet">
<link rel="icon" href="data:image/svg+xml,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 100 100'><circle cx='50' cy='50' r='48' fill='%23E60023'/><text x='50' y='72' font-size='60' text-anchor='middle' fill='white' font-family='serif' font-weight='bold'>P</text></svg>">
<style>
:root {
  --red: #E60023;
  --red-hover: #C4001E;
  --red-soft: rgba(230, 0, 35, 0.08);
  --blue: #3B82F6;
  --blue-soft: rgba(59, 130, 246, 0.10);
  --green: #10B981;
  --green-soft: rgba(16, 185, 129, 0.10);
  --amber: #F59E0B;
  --amber-soft: rgba(245, 158, 11, 0.10);
  --purple: #8B5CF6;
  --purple-soft: rgba(139, 92, 246, 0.10);
  --bg: #FAFAFB;
  --bg-elevated: #FFFFFF;
  --bg-sunken: #F4F4F6;
  --surface: #F1F1F4;
  --surface-2: #E8E8EC;
  --border: rgba(0,0,0,0.07);
  --border-strong: rgba(0,0,0,0.12);
  --text: #0A0A0B;
  --text-soft: #2A2A30;
  --text-muted: #6B6B76;
  --text-dim: #A0A0AA;
  --radius-sm: 10px;
  --radius: 14px;
  --radius-lg: 18px;
  --radius-xl: 24px;
  --radius-full: 999px;
  --shadow-xs: 0 1px 2px rgba(0,0,0,.04);
  --shadow-sm: 0 2px 8px rgba(0,0,0,.05), 0 1px 2px rgba(0,0,0,.04);
  --shadow: 0 8px 24px rgba(0,0,0,.07), 0 2px 6px rgba(0,0,0,.04);
  --shadow-lg: 0 20px 60px rgba(0,0,0,.12), 0 4px 12px rgba(0,0,0,.06);
  --shadow-red: 0 12px 32px rgba(230, 0, 35, 0.22);
  --ease: cubic-bezier(0.16, 1, 0.3, 1);
  --ease-spring: cubic-bezier(0.34, 1.56, 0.64, 1);
  --font: 'Inter', -apple-system, 'Segoe UI', Roboto, sans-serif;
  --mono: 'JetBrains Mono', 'SF Mono', monospace;
  --safe-top: env(safe-area-inset-top, 0px);
  --safe-bottom: env(safe-area-inset-bottom, 0px);
  --safe-left: env(safe-area-inset-left, 0px);
  --safe-right: env(safe-area-inset-right, 0px);
}
body[data-theme="dark"] {
  --bg: #08080A;
  --bg-elevated: #0F0F12;
  --bg-sunken: #050506;
  --surface: #17171B;
  --surface-2: #202026;
  --border: rgba(255, 255, 255, 0.07);
  --border-strong: rgba(255, 255, 255, 0.13);
  --text: #F7F7F9;
  --text-soft: #E5E5EA;
  --text-muted: #8E8E99;
  --text-dim: #57575F;
  --shadow-xs: 0 1px 2px rgba(0,0,0,.3);
  --shadow-sm: 0 2px 8px rgba(0,0,0,.35);
  --shadow: 0 8px 24px rgba(0,0,0,.45);
  --shadow-lg: 0 20px 60px rgba(0,0,0,.65);
}
* { box-sizing: border-box; margin: 0; padding: 0; -webkit-tap-highlight-color: transparent; }
html { scroll-behavior: smooth; scroll-padding-top: 90px; -webkit-text-size-adjust: 100%; }
body {
  font-family: var(--font);
  background: var(--bg);
  color: var(--text);
  min-height: 100vh;
  font-feature-settings: 'cv11', 'ss01';
  letter-spacing: -0.011em;
  -webkit-font-smoothing: antialiased;
  -moz-osx-font-smoothing: grayscale;
  transition: background .3s var(--ease), color .3s var(--ease);
}
a { color: inherit; }
button { font-family: inherit; cursor: pointer; }
::selection { background: var(--red); color: #fff; }

.topnav {
  position: sticky; top: 0; z-index: 100;
  background: color-mix(in srgb, var(--bg) 88%, transparent);
  backdrop-filter: saturate(180%) blur(20px);
  -webkit-backdrop-filter: saturate(180%) blur(20px);
  border-bottom: 1px solid var(--border);
  padding-top: var(--safe-top);
}
.topnav-inner {
  max-width: 1400px; margin: 0 auto;
  height: 64px;
  padding: 0 max(24px, var(--safe-left)) 0 max(24px, var(--safe-right));
  display: flex; align-items: center; gap: 18px;
}
.brand {
  display: flex; align-items: center; gap: 10px;
  text-decoration: none; flex-shrink: 0;
  padding: 4px; border-radius: var(--radius-full);
  transition: background .15s;
}
.brand:hover { background: var(--surface); }
.brand-badge {
  width: 34px; height: 34px; border-radius: 50%;
  background: linear-gradient(135deg, #FF1E3D 0%, var(--red) 50%, #8B0020 100%);
  color: #fff; display: grid; place-items: center;
  font-size: 1.1rem; font-weight: 800;
  box-shadow: var(--shadow-red), inset 0 1px 0 rgba(255,255,255,.25);
}
.brand-name { font-weight: 800; font-size: .98rem; letter-spacing: -0.03em; }
.brand-tag {
  font-size: .6rem; font-weight: 700;
  padding: 3px 8px; border-radius: 999px;
  background: var(--red); color: #fff;
  letter-spacing: .05em; text-transform: uppercase;
}
.topnav-spacer { flex: 1; }
.navlinks { display: flex; gap: 6px; align-items: center; flex-shrink: 0; }
.navlink {
  display: inline-flex; align-items: center; gap: 6px;
  padding: 8px 14px; border-radius: var(--radius-sm);
  font-size: .84rem; font-weight: 600;
  text-decoration: none; color: var(--text-soft);
  background: transparent; border: 1px solid transparent;
  transition: all .15s var(--ease);
}
.navlink:hover { background: var(--surface); border-color: var(--border); color: var(--text); }
.navlink.primary {
  background: var(--red); color: #fff; border-color: transparent;
  box-shadow: 0 2px 8px rgba(230,0,35,.25);
}
.navlink.primary:hover { background: var(--red-hover); transform: translateY(-1px); }
.theme-toggle {
  width: 38px; height: 38px; border-radius: 10px;
  border: 1px solid var(--border); background: var(--bg-elevated);
  color: var(--text-muted); display: grid; place-items: center;
  font-size: 1rem;
  transition: all .15s var(--ease);
}
.theme-toggle:hover { color: var(--text); border-color: var(--border-strong); }

.hero {
  position: relative;
  padding: 72px max(24px, var(--safe-left)) 56px max(24px, var(--safe-right));
  overflow: hidden;
  isolation: isolate;
}
.hero-bg {
  position: absolute; inset: 0;
  z-index: -1;
  background:
    radial-gradient(ellipse 900px 500px at 15% -20%, rgba(230,0,35,.09), transparent 60%),
    radial-gradient(ellipse 700px 400px at 85% 0%, rgba(139,92,246,.08), transparent 60%),
    radial-gradient(ellipse 600px 400px at 50% 110%, rgba(59,130,246,.05), transparent 60%);
}
body[data-theme="dark"] .hero-bg {
  background:
    radial-gradient(ellipse 900px 500px at 15% -20%, rgba(230,0,35,.16), transparent 60%),
    radial-gradient(ellipse 700px 400px at 85% 0%, rgba(139,92,246,.14), transparent 60%),
    radial-gradient(ellipse 600px 400px at 50% 110%, rgba(59,130,246,.10), transparent 60%);
}
.hero-grid {
  position: absolute; inset: 0; z-index: -1;
  background-image:
    linear-gradient(rgba(0,0,0,.03) 1px, transparent 1px),
    linear-gradient(90deg, rgba(0,0,0,.03) 1px, transparent 1px);
  background-size: 40px 40px;
  mask-image: radial-gradient(ellipse 800px 500px at 50% 30%, black, transparent 70%);
  -webkit-mask-image: radial-gradient(ellipse 800px 500px at 50% 30%, black, transparent 70%);
}
body[data-theme="dark"] .hero-grid {
  background-image:
    linear-gradient(rgba(255,255,255,.035) 1px, transparent 1px),
    linear-gradient(90deg, rgba(255,255,255,.035) 1px, transparent 1px);
}
.hero-inner { max-width: 1400px; margin: 0 auto; }
.hero-eyebrow {
  display: inline-flex; align-items: center; gap: 8px;
  padding: 6px 14px;
  background: var(--bg-elevated);
  border: 1px solid var(--border);
  border-radius: var(--radius-full);
  font-size: .78rem; font-weight: 600;
  color: var(--text-muted);
  margin-bottom: 22px;
  box-shadow: var(--shadow-xs);
}
.hero-eyebrow .pulse {
  width: 6px; height: 6px; border-radius: 50%;
  background: var(--green);
  box-shadow: 0 0 0 3px color-mix(in srgb, var(--green) 22%, transparent);
  animation: pulse 2s ease-in-out infinite;
}
@keyframes pulse { 0%, 100% { opacity: 1; transform: scale(1); } 50% { opacity: .55; transform: scale(.92); } }
.hero h1 {
  font-size: clamp(2.4rem, 5.5vw, 4rem);
  font-weight: 900;
  letter-spacing: -0.05em;
  line-height: 1.02;
  margin-bottom: 20px;
  background: linear-gradient(140deg, var(--text) 20%, var(--text-muted) 90%);
  -webkit-background-clip: text; background-clip: text;
  -webkit-text-fill-color: transparent;
  max-width: 800px;
}
.hero h1 .accent {
  background: linear-gradient(135deg, var(--red), #FF4D6A);
  -webkit-background-clip: text; background-clip: text;
  -webkit-text-fill-color: transparent;
}
.hero p {
  font-size: 1.08rem;
  color: var(--text-muted);
  line-height: 1.65;
  max-width: 640px;
  margin-bottom: 36px;
}
.hero-stats {
  display: flex; flex-wrap: wrap; gap: 12px;
  margin-top: 28px;
}
.stat-pill {
  display: flex; flex-direction: column; gap: 2px;
  padding: 14px 20px;
  background: var(--bg-elevated);
  border: 1px solid var(--border);
  border-radius: var(--radius);
  min-width: 130px;
  box-shadow: var(--shadow-xs);
  transition: all .2s var(--ease);
}
.stat-pill:hover { transform: translateY(-2px); box-shadow: var(--shadow-sm); border-color: var(--border-strong); }
.stat-pill-num {
  font-family: var(--mono);
  font-size: 1.5rem; font-weight: 700;
  letter-spacing: -0.03em;
  color: var(--text);
  font-variant-numeric: tabular-nums;
  line-height: 1.1;
}
.stat-pill-lbl {
  font-size: .72rem; font-weight: 600;
  color: var(--text-muted);
  text-transform: uppercase; letter-spacing: .08em;
  margin-top: 3px;
}

.layout {
  max-width: 1400px;
  margin: 0 auto;
  padding: 0 max(24px, var(--safe-left)) calc(80px + var(--safe-bottom)) max(24px, var(--safe-right));
  display: grid;
  grid-template-columns: 280px minmax(0, 1fr);
  gap: 48px;
  align-items: start;
}

.sidebar {
  position: sticky;
  top: calc(64px + var(--safe-top) + 24px);
  max-height: calc(100vh - 100px - var(--safe-top));
  display: flex; flex-direction: column;
  overflow: hidden;
}
.sidebar-search {
  display: flex; align-items: center; gap: 10px;
  height: 40px;
  padding: 0 14px;
  background: var(--surface);
  border: 1px solid transparent;
  border-radius: var(--radius-sm);
  margin-bottom: 16px;
  flex-shrink: 0;
  transition: all .15s;
}
.sidebar-search:focus-within {
  background: var(--bg-elevated);
  border-color: var(--red);
  box-shadow: 0 0 0 3px var(--red-soft);
}
.sidebar-search svg { color: var(--text-muted); flex-shrink: 0; }
.sidebar-search input {
  flex: 1; min-width: 0;
  border: none; outline: none; background: transparent;
  font: 500 .86rem var(--font);
  color: var(--text);
  height: 100%;
}
.sidebar-search input::placeholder { color: var(--text-dim); }
.sidebar-search .kbd {
  font-family: var(--mono); font-size: .66rem; font-weight: 600;
  padding: 2px 6px;
  background: var(--bg-elevated);
  border: 1px solid var(--border);
  border-radius: 5px;
  color: var(--text-dim);
  flex-shrink: 0;
}
.sidebar-nav {
  flex: 1; overflow-y: auto;
  padding-right: 8px;
  margin-right: -8px;
}
.nav-group { margin-bottom: 22px; }
.nav-group-title {
  display: flex; align-items: center; gap: 8px;
  padding: 0 8px 8px;
  font-size: .68rem; font-weight: 800;
  text-transform: uppercase; letter-spacing: .1em;
  color: var(--text-dim);
}
.nav-group-title .count {
  margin-left: auto;
  padding: 2px 7px;
  background: var(--surface);
  border-radius: 999px;
  font-size: .62rem;
  color: var(--text-muted);
  font-family: var(--mono);
}
.nav-item {
  display: flex; align-items: center; gap: 10px;
  padding: 8px 10px;
  border-radius: 8px;
  text-decoration: none;
  color: var(--text-soft);
  font-size: .84rem; font-weight: 500;
  transition: all .12s;
  cursor: pointer;
  border: 1px solid transparent;
}
.nav-item:hover { background: var(--surface); color: var(--text); }
.nav-item.active {
  background: var(--red-soft);
  color: var(--red);
  border-color: color-mix(in srgb, var(--red) 20%, transparent);
  font-weight: 600;
}
.nav-item .method {
  font-family: var(--mono);
  font-size: .58rem; font-weight: 700;
  padding: 2px 6px; border-radius: 5px;
  letter-spacing: .04em;
  flex-shrink: 0;
  min-width: 42px; text-align: center;
}
.method.get { background: var(--blue-soft); color: var(--blue); }
.method.post { background: var(--green-soft); color: var(--green); }
.method.put, .method.patch { background: var(--amber-soft); color: var(--amber); }
.method.delete { background: var(--red-soft); color: var(--red); }
.nav-item .path-text {
  font-family: var(--mono);
  font-size: .78rem;
  overflow: hidden; text-overflow: ellipsis; white-space: nowrap;
  flex: 1; min-width: 0;
}

.content { min-width: 0; padding-top: 8px; }
.content-header {
  display: flex; align-items: center; justify-content: space-between;
  gap: 16px; margin-bottom: 24px;
  padding-bottom: 20px;
  border-bottom: 1px solid var(--border);
  flex-wrap: wrap;
}
.content-header h2 {
  font-size: 1.6rem; font-weight: 800;
  letter-spacing: -0.035em;
}
.content-header .lead {
  color: var(--text-muted); font-size: .9rem;
  margin-top: 4px;
}
.filter-hint {
  font-size: .78rem; font-weight: 600;
  color: var(--text-muted);
  padding: 6px 12px;
  background: var(--surface);
  border-radius: var(--radius-full);
}

.endpoint {
  background: var(--bg-elevated);
  border: 1px solid var(--border);
  border-radius: var(--radius-lg);
  margin-bottom: 14px;
  overflow: hidden;
  transition: all .25s var(--ease);
  scroll-margin-top: 100px;
}
.endpoint:hover { border-color: var(--border-strong); box-shadow: var(--shadow-sm); }
.endpoint.highlight {
  border-color: color-mix(in srgb, var(--red) 40%, var(--border));
  box-shadow: 0 0 0 3px var(--red-soft);
}
.endpoint-head {
  display: flex; align-items: center; gap: 14px;
  padding: 16px 20px;
  cursor: pointer;
  user-select: none;
  transition: background .15s;
}
.endpoint-head:hover { background: var(--surface); }
.endpoint.open .endpoint-head { background: var(--surface); border-bottom: 1px solid var(--border); }
.endpoint-method {
  font-family: var(--mono);
  font-size: .68rem; font-weight: 700;
  padding: 6px 11px; border-radius: 8px;
  letter-spacing: .05em;
  flex-shrink: 0;
  min-width: 62px; text-align: center;
}
.endpoint-method.get { background: var(--blue); color: #fff; box-shadow: 0 2px 8px rgba(59,130,246,.3); }
.endpoint-method.post { background: var(--green); color: #fff; box-shadow: 0 2px 8px rgba(16,185,129,.3); }
.endpoint-method.put, .endpoint-method.patch { background: var(--amber); color: #fff; box-shadow: 0 2px 8px rgba(245,158,11,.3); }
.endpoint-method.delete { background: var(--red); color: #fff; box-shadow: 0 2px 8px rgba(230,0,35,.3); }
.endpoint-path {
  font-family: var(--mono);
  font-size: .92rem; font-weight: 600;
  color: var(--text);
  flex-shrink: 0;
  white-space: nowrap; overflow: hidden; text-overflow: ellipsis;
  min-width: 0;
}
.endpoint-summary {
  color: var(--text-muted);
  font-size: .86rem;
  overflow: hidden; text-overflow: ellipsis; white-space: nowrap;
  flex: 1; min-width: 0;
  padding-left: 12px;
  border-left: 1px solid var(--border);
  margin-left: 4px;
}
.endpoint-chevron {
  color: var(--text-dim); font-size: .9rem;
  flex-shrink: 0;
  transition: transform .25s var(--ease);
}
.endpoint.open .endpoint-chevron { transform: rotate(90deg); }
.endpoint-body {
  display: none;
  padding: 22px 24px 26px;
}
.endpoint.open .endpoint-body { display: block; animation: slideDown .3s var(--ease); }
@keyframes slideDown {
  from { opacity: 0; transform: translateY(-6px); }
  to { opacity: 1; transform: translateY(0); }
}
.endpoint-desc {
  color: var(--text-soft);
  font-size: .9rem; line-height: 1.7;
  margin-bottom: 20px;
}
.endpoint-section { margin-bottom: 22px; }
.endpoint-section:last-child { margin-bottom: 0; }
.section-title {
  display: flex; align-items: center; gap: 8px;
  font-size: .7rem; font-weight: 800;
  text-transform: uppercase; letter-spacing: .1em;
  color: var(--text-dim);
  margin-bottom: 12px;
}
.section-title::after {
  content: ''; flex: 1; height: 1px;
  background: var(--border);
}
.param-table {
  border: 1px solid var(--border);
  border-radius: var(--radius);
  overflow: hidden;
}
.param-row {
  display: grid;
  grid-template-columns: minmax(140px, 1fr) minmax(80px, auto) minmax(60px, auto) 2fr;
  gap: 16px;
  padding: 12px 16px;
  border-bottom: 1px solid var(--border);
  align-items: baseline;
  font-size: .84rem;
}
.param-row:last-child { border-bottom: none; }
.param-row.header {
  background: var(--bg-sunken);
  font-size: .68rem; font-weight: 800;
  text-transform: uppercase; letter-spacing: .08em;
  color: var(--text-dim);
  padding: 10px 16px;
}
.param-name {
  font-family: var(--mono); font-weight: 600;
  color: var(--text);
  overflow-wrap: break-word;
}
.param-type {
  font-family: var(--mono); font-size: .78rem;
  color: var(--purple);
  font-weight: 500;
}
.param-required {
  font-size: .68rem; font-weight: 700;
  padding: 2px 8px; border-radius: 999px;
  background: var(--red-soft); color: var(--red);
  text-align: center;
  text-transform: uppercase; letter-spacing: .04em;
}
.param-required.optional {
  background: var(--surface); color: var(--text-dim);
}
.param-desc { color: var(--text-muted); font-size: .82rem; }

.code-block {
  position: relative;
  background: #0B0B0F;
  border-radius: var(--radius);
  padding: 18px 20px;
  font-family: var(--mono);
  font-size: .82rem;
  color: #E5E5EA;
  overflow-x: auto;
  line-height: 1.65;
  border: 1px solid rgba(255,255,255,.05);
}
body[data-theme="dark"] .code-block { background: #050507; }
.code-block pre { margin: 0; white-space: pre; }
.copy-btn {
  position: absolute; top: 10px; right: 10px;
  display: inline-flex; align-items: center; gap: 6px;
  padding: 5px 10px;
  background: rgba(255,255,255,.08);
  color: rgba(255,255,255,.7);
  border: 1px solid rgba(255,255,255,.1);
  border-radius: 7px;
  font-size: .7rem; font-weight: 600;
  transition: all .15s;
  font-family: var(--font);
}
.copy-btn:hover { background: rgba(255,255,255,.15); color: #fff; }
.copy-btn.copied { background: var(--green); color: #fff; border-color: transparent; }
.copy-btn svg { width: 12px; height: 12px; }

.tok-key { color: #7EE8FA; }
.tok-string { color: #A5F3A0; }
.tok-num { color: #FFB86C; }
.tok-bool { color: #FF79C6; }
.tok-null { color: #8E8E99; }
.tok-punc { color: #6B6B76; }

.response-row {
  display: flex; align-items: center; gap: 12px;
  padding: 10px 0;
  border-bottom: 1px solid var(--border);
}
.response-row:last-child { border-bottom: none; }
.status-code {
  font-family: var(--mono);
  font-size: .78rem; font-weight: 700;
  padding: 3px 10px; border-radius: 6px;
  min-width: 52px; text-align: center;
}
.status-code.ok { background: var(--green-soft); color: var(--green); }
.status-code.redirect { background: var(--blue-soft); color: var(--blue); }
.status-code.client { background: var(--amber-soft); color: var(--amber); }
.status-code.server { background: var(--red-soft); color: var(--red); }
.response-desc { color: var(--text-muted); font-size: .84rem; }

.try-it {
  display: flex; align-items: center; gap: 12px;
  padding-top: 16px;
  border-top: 1px solid var(--border);
  flex-wrap: wrap;
}
.btn {
  display: inline-flex; align-items: center; gap: 8px;
  padding: 10px 18px;
  border: 1px solid transparent;
  border-radius: var(--radius-sm);
  font-size: .84rem; font-weight: 600;
  transition: all .15s var(--ease);
  text-decoration: none;
  cursor: pointer;
  font-family: inherit;
  min-height: 40px;
}
.btn:active { transform: scale(.97); }
.btn.primary {
  background: var(--red); color: #fff;
  box-shadow: 0 2px 8px rgba(230,0,35,.25);
}
.btn.primary:hover { background: var(--red-hover); transform: translateY(-1px); box-shadow: 0 6px 18px rgba(230,0,35,.35); }
.btn.ghost {
  background: var(--bg-elevated); color: var(--text-soft);
  border-color: var(--border);
}
.btn.ghost:hover { background: var(--surface); border-color: var(--border-strong); }
.btn:disabled { opacity: .5; cursor: not-allowed; }
.try-status {
  font-family: var(--mono); font-size: .78rem;
  color: var(--text-muted);
  margin-left: auto;
}
.try-status.ok { color: var(--green); font-weight: 700; }
.try-status.err { color: var(--red); font-weight: 700; }

.loader {
  width: 14px; height: 14px;
  border: 2px solid rgba(255,255,255,.3);
  border-top-color: #fff;
  border-radius: 50%;
  animation: spin .7s linear infinite;
}
@keyframes spin { to { transform: rotate(360deg); } }

.empty-state {
  text-align: center;
  padding: 60px 20px;
  color: var(--text-muted);
}
.empty-state svg { color: var(--text-dim); margin-bottom: 14px; }
.empty-state h3 { color: var(--text); font-weight: 700; margin-bottom: 6px; font-size: 1rem; }

.docs-footer {
  border-top: 1px solid var(--border);
  background: var(--bg-sunken);
  padding: 32px max(24px, var(--safe-left)) calc(32px + var(--safe-bottom)) max(24px, var(--safe-right));
  margin-top: 60px;
}
.docs-footer-inner {
  max-width: 1400px; margin: 0 auto;
  display: flex; align-items: center; justify-content: space-between;
  flex-wrap: wrap; gap: 16px;
  font-size: .84rem; color: var(--text-muted);
}
.docs-footer-inner .heart { color: var(--red); }

.toast-container {
  position: fixed;
  bottom: calc(24px + var(--safe-bottom));
  right: max(24px, var(--safe-right));
  display: flex; flex-direction: column; gap: 8px;
  z-index: 500;
  pointer-events: none;
}
.toast {
  padding: 12px 18px;
  background: var(--bg-elevated);
  border: 1px solid var(--border-strong);
  border-radius: var(--radius-sm);
  box-shadow: var(--shadow-lg);
  font-size: .84rem; font-weight: 500;
  color: var(--text);
  pointer-events: auto;
  animation: toastIn .3s var(--ease-spring);
  display: flex; align-items: center; gap: 10px;
}
@keyframes toastIn { from { transform: translateX(120%); opacity: 0; } to { transform: translateX(0); opacity: 1; } }
.toast .dot { width: 6px; height: 6px; border-radius: 50%; background: var(--green); }

::-webkit-scrollbar { width: 10px; height: 10px; }
::-webkit-scrollbar-track { background: transparent; }
::-webkit-scrollbar-thumb { background: var(--surface-2); border-radius: 10px; border: 2px solid transparent; background-clip: padding-box; }
::-webkit-scrollbar-thumb:hover { background-color: var(--text-dim); background-clip: padding-box; }
* { scrollbar-width: thin; scrollbar-color: var(--surface-2) transparent; }

@media (max-width: 1024px) {
  .layout { grid-template-columns: 240px minmax(0, 1fr); gap: 32px; }
}
@media (max-width: 860px) {
  .hero { padding: 48px 20px 40px; }
  .hero h1 { font-size: 2.2rem; }
  .layout {
    grid-template-columns: 1fr;
    gap: 0;
    padding: 0 16px calc(60px + var(--safe-bottom));
  }
  .sidebar {
    position: static;
    max-height: none;
    margin-bottom: 24px;
  }
  .sidebar-nav {
    display: flex; gap: 8px; overflow-x: auto; padding-bottom: 8px;
    margin-right: 0; padding-right: 0;
  }
  .nav-group { margin-bottom: 0; flex-shrink: 0; }
  .nav-group-title { display: none; }
  .nav-group > div { display: flex; gap: 6px; }
  .nav-item { flex-shrink: 0; white-space: nowrap; }
  .param-row { grid-template-columns: 1fr; gap: 6px; }
  .param-row.header { display: none; }
  .param-name::before { content: 'Name: '; color: var(--text-dim); font-family: var(--font); font-size: .72rem; font-weight: 600; text-transform: uppercase; letter-spacing: .06em; margin-right: 6px; }
  .endpoint-summary { display: none; }
  .endpoint-path { font-size: .82rem; }
  .brand-name, .brand-tag { display: none; }
  .topnav-inner { padding: 0 16px; height: 58px; }
  .navlink span { display: none; }
  .navlink { padding: 8px 10px; }
}
@media (max-width: 480px) {
  .hero h1 { font-size: 1.8rem; }
  .hero p { font-size: .96rem; }
  .stat-pill { min-width: 100px; padding: 10px 14px; }
  .stat-pill-num { font-size: 1.15rem; }
  .content-header h2 { font-size: 1.3rem; }
}
@media (prefers-reduced-motion: reduce) {
  *, *::before, *::after {
    animation-duration: .01ms !important;
    animation-iteration-count: 1 !important;
    transition-duration: .01ms !important;
  }
}
</style>
</head>
<body>

<header class="topnav">
  <div class="topnav-inner">
    <a class="brand" href="/">
      <span class="brand-badge">P</span>
      <span class="brand-name">Pinterest Scraper</span>
      <span class="brand-tag">API</span>
    </a>
    <div class="topnav-spacer"></div>
    <nav class="navlinks">
      <a class="navlink" href="/" title="Home">
        <svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><path d="M3 9l9-7 9 7v11a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2z"/><polyline points="9 22 9 12 15 12 15 22"/></svg>
        <span>Home</span>
      </a>
      <a class="navlink" href="/api/health" target="_blank" title="Health">
        <svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><path d="M22 12h-4l-3 9L9 3l-3 9H2"/></svg>
        <span>Health</span>
      </a>
      <a class="navlink primary" href="/openapi.json" target="_blank" title="OpenAPI JSON">
        <svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><polyline points="16 18 22 12 16 6"/><polyline points="8 6 2 12 8 18"/></svg>
        <span>OpenAPI</span>
      </a>
      <button class="theme-toggle" id="theme-toggle" aria-label="Toggle theme">
        <svg id="theme-icon" width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><path d="M21 12.79A9 9 0 1 1 11.21 3 7 7 0 0 0 21 12.79z"/></svg>
      </button>
    </nav>
  </div>
</header>

<section class="hero">
  <div class="hero-bg"></div>
  <div class="hero-grid"></div>
  <div class="hero-inner">
    <div class="hero-eyebrow">
      <span class="pulse"></span>
      <span id="hero-version">Loading spec…</span>
    </div>
    <h1>Pinterest Scraper <span class="accent">API</span></h1>
    <p id="hero-desc">High-quality Pinterest scraper API with job events, exports, and gallery management.</p>
    <div class="hero-stats" id="hero-stats"></div>
  </div>
</section>

<div class="layout">
  <aside class="sidebar">
    <label class="sidebar-search">
      <svg width="15" height="15" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><circle cx="11" cy="11" r="8"/><line x1="21" y1="21" x2="16.65" y2="16.65"/></svg>
      <input id="filter-input" type="text" placeholder="Filter endpoints…" autocomplete="off">
      <span class="kbd">/</span>
    </label>
    <nav class="sidebar-nav" id="sidebar-nav"></nav>
  </aside>

  <main class="content">
    <div class="content-header">
      <div>
        <h2>API Reference</h2>
        <div class="lead" id="content-lead">All endpoints exposed by the Pinterest Scraper service.</div>
      </div>
      <div class="filter-hint" id="filter-hint" style="display:none"></div>
    </div>
    <div id="endpoints-container"></div>
  </main>
</div>

<footer class="docs-footer">
  <div class="docs-footer-inner">
    <div>Made with <span class="heart">♥</span> using FastAPI</div>
    <div id="footer-meta"></div>
  </div>
</footer>

<div class="toast-container" id="toast-container"></div>

<script>
(function () {
  const SPEC_URL = '/openapi.json';

  const METHOD_ORDER = ['get', 'post', 'put', 'patch', 'delete'];
  const METHOD_LABEL = { get: 'GET', post: 'POST', put: 'PUT', patch: 'PATCH', delete: 'DELETE' };

  const $ = (id) => document.getElementById(id);

  function escapeHtml(s) {
    return String(s == null ? '' : s).replace(/[&<>"']/g, c => ({
      '&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;', "'": '&#39;'
    }[c]));
  }

  function toast(msg) {
    const el = document.createElement('div');
    el.className = 'toast';
    el.innerHTML = '<span class="dot"></span><span>' + escapeHtml(msg) + '</span>';
    $('toast-container').appendChild(el);
    setTimeout(() => {
      el.style.transition = 'opacity .25s, transform .25s';
      el.style.opacity = '0';
      el.style.transform = 'translateX(120%)';
      setTimeout(() => el.remove(), 260);
    }, 2000);
  }

  function initTheme() {
    const saved = localStorage.getItem('docs-theme') || localStorage.getItem('theme');
    const prefers = window.matchMedia('(prefers-color-scheme: dark)').matches ? 'dark' : 'light';
    document.body.dataset.theme = saved || prefers;
    updateIcon();
    $('theme-toggle').addEventListener('click', () => {
      const next = document.body.dataset.theme === 'dark' ? 'light' : 'dark';
      document.body.dataset.theme = next;
      try { localStorage.setItem('docs-theme', next); } catch (e) {}
      updateIcon();
    });
  }

  function updateIcon() {
    const dark = document.body.dataset.theme === 'dark';
    $('theme-icon').innerHTML = dark
      ? '<circle cx="12" cy="12" r="4"/><path d="M12 2v2M12 20v2M4.93 4.93l1.41 1.41M17.66 17.66l1.41 1.41M2 12h2M20 12h2M4.93 19.07l1.41-1.41M17.66 6.34l1.41-1.41"/>'
      : '<path d="M21 12.79A9 9 0 1 1 11.21 3 7 7 0 0 0 21 12.79z"/>';
  }

  function resolveRef(ref, spec) {
    if (!ref || !ref.startsWith('#/')) return null;
    const parts = ref.slice(2).split('/');
    let cur = spec;
    for (const p of parts) {
      if (cur == null) return null;
      cur = cur[p.replace(/~1/g, '/').replace(/~0/g, '~')];
    }
    return cur;
  }

  function schemaType(schema) {
    if (!schema) return 'any';
    if (schema.$ref) return schema.$ref.split('/').pop();
    if (schema.anyOf || schema.oneOf) {
      const arr = schema.anyOf || schema.oneOf;
      return arr.map(schemaType).filter(Boolean).join(' | ');
    }
    if (schema.type === 'array') {
      return 'array<' + schemaType(schema.items) + '>';
    }
    return schema.type || 'object';
  }

  function stringifySample(value, indent, depth) {
    indent = indent || 0;
    depth = depth || 0;
    if (depth > 6) return 'null';
    const pad = '  '.repeat(indent);
    const padIn = '  '.repeat(indent + 1);
    if (value === null) return '<span class="tok-null">null</span>';
    if (typeof value === 'boolean') return '<span class="tok-bool">' + value + '</span>';
    if (typeof value === 'number') return '<span class="tok-num">' + value + '</span>';
    if (typeof value === 'string') return '<span class="tok-string">"' + escapeHtml(value) + '"</span>';
    if (Array.isArray(value)) {
      if (value.length === 0) return '<span class="tok-punc">[]</span>';
      return '[\n' + value.map(v => padIn + stringifySample(v, indent + 1, depth + 1)).join(',\n') + '\n' + pad + ']';
    }
    if (typeof value === 'object') {
      const keys = Object.keys(value);
      if (keys.length === 0) return '<span class="tok-punc">{}</span>';
      return '{\n' + keys.map(k =>
        padIn + '<span class="tok-key">"' + escapeHtml(k) + '"</span><span class="tok-punc">:</span> ' +
        stringifySample(value[k], indent + 1, depth + 1)
      ).join(',\n') + '\n' + pad + '}';
    }
    return 'null';
  }

  function sampleFromSchema(schema, spec, depth) {
    depth = depth || 0;
    if (!schema || depth > 5) return null;
    if (schema.$ref) {
      const resolved = resolveRef(schema.$ref, spec);
      return sampleFromSchema(resolved, spec, depth + 1);
    }
    if (schema.example !== undefined) return schema.example;
    if (schema.default !== undefined) return schema.default;
    if (schema.anyOf || schema.oneOf) {
      return sampleFromSchema((schema.anyOf || schema.oneOf)[0], spec, depth + 1);
    }
    const t = schema.type;
    if (t === 'object' || schema.properties) {
      const out = {};
      const props = schema.properties || {};
      for (const k of Object.keys(props)) {
        out[k] = sampleFromSchema(props[k], spec, depth + 1);
      }
      return out;
    }
    if (t === 'array') {
      return [sampleFromSchema(schema.items, spec, depth + 1)];
    }
    if (t === 'string') {
      if (schema.enum && schema.enum.length) return schema.enum[0];
      if (schema.format === 'date-time') return new Date().toISOString();
      return 'string';
    }
    if (t === 'integer') return 0;
    if (t === 'number') return 0;
    if (t === 'boolean') return false;
    return null;
  }

  function buildEndpointData(path, method, op, spec) {
    const params = [];
    const allParams = [].concat(op.parameters || []);
    if (op.requestBody) {
      const rb = op.requestBody;
      const content = (rb.content && (rb.content['application/json'] || Object.values(rb.content)[0])) || null;
      const schema = content && content.schema;
      const resolved = schema && schema.$ref ? resolveRef(schema.$ref, spec) : schema;
      const props = (resolved && resolved.properties) || {};
      for (const name of Object.keys(props)) {
        params.push({
          name: name,
          type: schemaType(props[name]),
          required: (resolved.required || []).includes(name),
          desc: props[name].description || '',
          in: 'body',
        });
      }
    }
    for (const p of allParams) {
      params.push({
        name: p.name,
        type: schemaType(p.schema || {}),
        required: !!p.required,
        desc: p.description || '',
        in: p.in || 'query',
      });
    }

    let exampleJSON = null;
    if (op.requestBody) {
      const content = op.requestBody.content || {};
      const key = content['application/json'] ? 'application/json' : Object.keys(content)[0];
      if (key) {
        const schema = content[key].schema;
        const resolved = schema && schema.$ref ? resolveRef(schema.$ref, spec) : schema;
        exampleJSON = sampleFromSchema(resolved, spec, 0);
      }
    }

    return {
      path: path,
      method: method,
      methodLabel: METHOD_LABEL[method] || method.toUpperCase(),
      summary: op.summary || '',
      description: op.description || '',
      tags: op.tags || ['default'],
      params: params,
      exampleJSON: exampleJSON,
      responses: op.responses || {},
    };
  }

  function renderHero(spec, endpoints) {
    $('hero-version').textContent = 'v' + (spec.info.version || '1.0.0') + ' · Live';
    $('hero-desc').textContent = spec.info.description || '';
    const methods = new Set(endpoints.map(e => e.method));
    const tags = new Set(endpoints.map(e => e.tags[0]));
    const stats = [
      { num: endpoints.length, lbl: 'Endpoints' },
      { num: tags.size, lbl: 'Groups' },
      { num: methods.size, lbl: 'HTTP methods' },
    ];
    $('hero-stats').innerHTML = stats.map(s =>
      '<div class="stat-pill"><div class="stat-pill-num">' + s.num + '</div><div class="stat-pill-lbl">' + s.lbl + '</div></div>'
    ).join('');
    $('footer-meta').textContent = 'API v' + (spec.info.version || '1.0.0');
  }

  function renderSidebar(endpoints) {
    const nav = $('sidebar-nav');
    const groups = {};
    for (const ep of endpoints) {
      const tag = ep.tags[0] || 'default';
      (groups[tag] = groups[tag] || []).push(ep);
    }
    const tagOrder = Object.keys(groups).sort();
    nav.innerHTML = tagOrder.map(tag => {
      const items = groups[tag].map(ep => {
        const id = 'ep-' + ep.method + '-' + ep.path.replace(/[^\w]/g, '_');
        return '<a class="nav-item" href="#' + id + '" data-path="' + escapeHtml(ep.path) + '">' +
          '<span class="method ' + ep.method + '">' + ep.methodLabel + '</span>' +
          '<span class="path-text">' + escapeHtml(ep.path) + '</span>' +
        '</a>';
      }).join('');
      return '<div class="nav-group">' +
        '<div class="nav-group-title">' + escapeHtml(tag) + '<span class="count">' + groups[tag].length + '</span></div>' +
        '<div>' + items + '</div>' +
      '</div>';
    }).join('');
  }

  function renderEndpoint(ep) {
    const id = 'ep-' + ep.method + '-' + ep.path.replace(/[^\w]/g, '_');
    const paramsHTML = ep.params.length ? (
      '<div class="endpoint-section">' +
        '<div class="section-title">Parameters</div>' +
        '<div class="param-table">' +
          '<div class="param-row header"><div>Name</div><div>Type</div><div>Required</div><div>Description</div></div>' +
          ep.params.map(p =>
            '<div class="param-row">' +
              '<div class="param-name">' + escapeHtml(p.name) + '</div>' +
              '<div class="param-type">' + escapeHtml(p.type) + '</div>' +
              '<div class="param-required' + (p.required ? '' : ' optional') + '">' + (p.required ? 'Required' : 'Optional') + '</div>' +
              '<div class="param-desc">' + escapeHtml(p.desc) + '</div>' +
            '</div>'
          ).join('') +
        '</div>' +
      '</div>'
    ) : '';

    const bodyHTML = ep.exampleJSON != null ? (
      '<div class="endpoint-section">' +
        '<div class="section-title">Request body</div>' +
        '<div class="code-block">' +
          '<button class="copy-btn" data-copy="' + escapeHtml(JSON.stringify(ep.exampleJSON, null, 2)) + '">' +
            '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.4" stroke-linecap="round" stroke-linejoin="round"><rect x="9" y="9" width="13" height="13" rx="2" ry="2"/><path d="M5 15H4a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2h9a2 2 0 0 1 2 2v1"/></svg>' +
            '<span>Copy</span>' +
          '</button>' +
          '<pre>' + stringifySample(ep.exampleJSON, 0, 0) + '</pre>' +
        '</div>' +
      '</div>'
    ) : '';

    const responsesHTML = Object.keys(ep.responses).length ? (
      '<div class="endpoint-section">' +
        '<div class="section-title">Responses</div>' +
        Object.entries(ep.responses).map(([code, resp]) => {
          let cls = 'ok';
          const n = parseInt(code, 10);
          if (n >= 500) cls = 'server';
          else if (n >= 400) cls = 'client';
          else if (n >= 300) cls = 'redirect';
          const desc = (resp && (resp.description || '')) || '';
          return '<div class="response-row">' +
            '<span class="status-code ' + cls + '">' + escapeHtml(code) + '</span>' +
            '<span class="response-desc">' + escapeHtml(desc) + '</span>' +
          '</div>';
        }).join('') +
      '</div>'
    ) : '';

    const tryHTML =
      '<div class="try-it">' +
        '<button class="btn primary" data-try="' + escapeHtml(ep.method) + '" data-path="' + escapeHtml(ep.path) + '">' +
          '<svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.4" stroke-linecap="round" stroke-linejoin="round"><polygon points="5 3 19 12 5 21 5 3"/></svg>' +
          '<span>Try it</span>' +
        '</button>' +
        '<button class="btn ghost" data-copy="' + escapeHtml(window.location.origin + ep.path) + '">' +
          '<svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.4" stroke-linecap="round" stroke-linejoin="round"><rect x="9" y="9" width="13" height="13" rx="2" ry="2"/><path d="M5 15H4a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2h9a2 2 0 0 1 2 2v1"/></svg>' +
          '<span>Copy URL</span>' +
        '</button>' +
        '<span class="try-status"></span>' +
      '</div>';

    return '<article class="endpoint" id="' + id + '">' +
      '<div class="endpoint-head" data-toggle>' +
        '<span class="endpoint-method ' + ep.method + '">' + ep.methodLabel + '</span>' +
        '<span class="endpoint-path">' + escapeHtml(ep.path) + '</span>' +
        '<span class="endpoint-summary">' + escapeHtml(ep.summary) + '</span>' +
        '<svg class="endpoint-chevron" width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.4" stroke-linecap="round" stroke-linejoin="round"><polyline points="9 18 15 12 9 6"/></svg>' +
      '</div>' +
      '<div class="endpoint-body">' +
        (ep.description ? '<div class="endpoint-desc">' + escapeHtml(ep.description) + '</div>' : '') +
        paramsHTML +
        bodyHTML +
        responsesHTML +
        tryHTML +
      '</div>' +
    '</article>';
  }

  function renderContent(endpoints) {
    $('endpoints-container').innerHTML = endpoints.map(renderEndpoint).join('');
  }

  function wireInteractions(spec) {
    document.addEventListener('click', (e) => {
      const toggle = e.target.closest('[data-toggle]');
      if (toggle) {
        const ep = toggle.closest('.endpoint');
        if (ep) ep.classList.toggle('open');
        return;
      }
      const copyBtn = e.target.closest('.copy-btn');
      if (copyBtn) {
        const text = copyBtn.getAttribute('data-copy') || '';
        navigator.clipboard.writeText(text).then(() => {
          const old = copyBtn.innerHTML;
          copyBtn.classList.add('copied');
          copyBtn.innerHTML = '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="3" stroke-linecap="round" stroke-linejoin="round"><polyline points="20 6 9 17 4 12"/></svg><span>Copied</span>';
          setTimeout(() => {
            copyBtn.innerHTML = old;
            copyBtn.classList.remove('copied');
          }, 1400);
          toast('Copied to clipboard');
        }).catch(() => {});
        return;
      }
      const tryBtn = e.target.closest('[data-try]');
      if (tryBtn) {
        const method = tryBtn.getAttribute('data-try');
        const path = tryBtn.getAttribute('data-path');
        const status = tryBtn.parentElement.querySelector('.try-status');
        tryBtn.disabled = true;
        const oldHtml = tryBtn.innerHTML;
        tryBtn.innerHTML = '<span class="loader"></span><span>Sending…</span>';
        status.className = 'try-status';
        status.textContent = '';
        const opts = { method: method.toUpperCase(), headers: { 'Accept': 'application/json' } };
        fetch(path, opts).then(async (res) => {
          const text = await res.text();
          if (res.ok) {
            status.className = 'try-status ok';
            status.textContent = res.status + ' ' + res.statusText;
          } else {
            status.className = 'try-status err';
            status.textContent = res.status + ' ' + res.statusText;
          }
        }).catch(err => {
          status.className = 'try-status err';
          status.textContent = 'Network error';
        }).finally(() => {
          tryBtn.disabled = false;
          tryBtn.innerHTML = oldHtml;
        });
        return;
      }
    });

    document.addEventListener('keydown', (e) => {
      const isInput = ['INPUT', 'TEXTAREA'].includes(document.activeElement.tagName);
      if (e.key === '/' && !isInput) {
        e.preventDefault();
        $('filter-input').focus();
      } else if (e.key === 'Escape' && document.activeElement === $('filter-input')) {
        $('filter-input').value = '';
        applyFilter('');
        $('filter-input').blur();
      }
    });

    const filterInput = $('filter-input');
    filterInput.addEventListener('input', () => applyFilter(filterInput.value));
  }

  let _allEndpoints = [];

  function applyFilter(q) {
    const term = (q || '').trim().toLowerCase();
    const filtered = term
      ? _allEndpoints.filter(ep =>
          ep.path.toLowerCase().includes(term) ||
          (ep.summary || '').toLowerCase().includes(term) ||
          (ep.description || '').toLowerCase().includes(term) ||
          ep.tags.some(t => t.toLowerCase().includes(term)))
      : _allEndpoints;

    renderContent(filtered);
    renderSidebar(filtered);

    const hint = $('filter-hint');
    if (term) {
      hint.style.display = '';
      hint.textContent = filtered.length + ' match' + (filtered.length === 1 ? '' : 'es') + ' for "' + term + '"';
    } else {
      hint.style.display = 'none';
    }

    const container = $('endpoints-container');
    if (!filtered.length) {
      container.innerHTML = '<div class="empty-state">' +
        '<svg width="40" height="40" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.8" stroke-linecap="round" stroke-linejoin="round"><circle cx="11" cy="11" r="8"/><line x1="21" y1="21" x2="16.65" y2="16.65"/></svg>' +
        '<h3>No endpoints found</h3>' +
        '<p>Try a different search term.</p>' +
      '</div>';
    }
  }

  function onScrollHighlight() {
    const articles = Array.from(document.querySelectorAll('.endpoint'));
    if (!articles.length) return;
    let active = null;
    const top = 130;
    for (const a of articles) {
      const rect = a.getBoundingClientRect();
      if (rect.top <= top) active = a.id;
      else break;
    }
    document.querySelectorAll('.nav-item').forEach(item => {
      const isActive = active && item.getAttribute('href') === '#' + active;
      item.classList.toggle('active', !!isActive);
    });
  }

  function boot(spec) {
    const endpoints = [];
    const paths = spec.paths || {};
    for (const path of Object.keys(paths)) {
      const methods = paths[path];
      for (const method of METHOD_ORDER) {
        if (methods[method]) {
          endpoints.push(buildEndpointData(path, method, methods[method], spec));
        }
      }
    }
    endpoints.sort((a, b) => {
      const t = a.tags[0].localeCompare(b.tags[0]);
      if (t !== 0) return t;
      return a.path.localeCompare(b.path);
    });

    _allEndpoints = endpoints;
    renderHero(spec, endpoints);
    renderSidebar(endpoints);
    renderContent(endpoints);
    wireInteractions(spec);

    window.addEventListener('scroll', onScrollHighlight, { passive: true });
    setTimeout(onScrollHighlight, 100);

    document.querySelectorAll('.nav-item').forEach(item => {
      item.addEventListener('click', () => {
        const id = item.getAttribute('href').slice(1);
        const el = document.getElementById(id);
        if (el) {
          el.classList.add('open', 'highlight');
          setTimeout(() => el.classList.remove('highlight'), 1500);
        }
      });
    });
  }

  function fail(err) {
    $('hero-version').textContent = 'Failed to load spec';
    $('hero-desc').textContent = String(err && err.message || err);
    $('endpoints-container').innerHTML =
      '<div class="empty-state">' +
        '<h3>Could not load OpenAPI spec</h3>' +
        '<p>Check that /openapi.json is reachable.</p>' +
      '</div>';
  }

  initTheme();
  fetch(SPEC_URL).then(r => r.json()).then(boot).catch(fail);
})();
</script>
</body>
</html>
"""


@app.get("/docs", include_in_schema=False)
async def custom_docs():
    return HTMLResponse(DOCS_HTML)


@app.get("/redoc", include_in_schema=False)
async def custom_redoc():
    from fastapi.openapi.docs import get_redoc_html
    return get_redoc_html(
        openapi_url=app.openapi_url,
        title="Pinterest Scraper · ReDoc",
    )


@app.get("/")
def index():
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/api/health")
def health():
    return {"ok": True}


@app.post("/api/scrape")
def start_scrape(req: ScrapeRequest):
    job = Job(req, _out_dir())
    JOBS[job.id] = job
    threading.Thread(target=job.run, daemon=True).start()
    return {"job_id": job.id}


@app.get("/api/jobs/{job_id}/events")
async def job_events(job_id: str):
    job = JOBS.get(job_id)
    if not job:
        raise HTTPException(404, "unknown job")

    async def gen():
        loop = asyncio.get_running_loop()
        while True:
            try:
                ev = await asyncio.wait_for(
                    loop.run_in_executor(None, job.events.get, True, 0.2), 5
                )
                yield f"data: {ev}\n\n"
                if json.loads(ev).get("event") == "done":
                    return
            except (queue.Empty, asyncio.TimeoutError):
                if job.done.is_set() and job.events.empty():
                    return
                yield ": keepalive\n\n"

    return StreamingResponse(gen(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache"})


@app.get("/api/jobs/{job_id}/result")
def job_result(job_id: str):
    job = JOBS.get(job_id)
    if not job:
        raise HTTPException(404, "unknown job")
    if not job.done.is_set():
        raise HTTPException(409, "job still running")
    return {"status": job.status, "stats": job.stats, "error": job.error, "pins": job.pins}


@app.post("/api/jobs/{job_id}/cancel")
def cancel_job(job_id: str):
    job = JOBS.get(job_id)
    if not job:
        raise HTTPException(404, "unknown job")
    job.cancelled = True
    job.status = "cancelled"
    job.done.set()
    job.emit(event="done", status="cancelled", total=len(job.pins), stats=job.stats)
    return {"ok": True}


@app.get("/api/jobs/{job_id}/images/{name}")
def job_image(job_id: str, name: str):
    job = JOBS.get(job_id)
    if not job:
        raise HTTPException(404, "unknown job")
    path = (job.out_dir / "images" / name).resolve()
    if not str(path).startswith(str(job.out_dir.resolve())) or not path.is_file():
        raise HTTPException(404, "not found")
    return FileResponse(path)


@app.get("/api/jobs/{job_id}/export/{fmt}")
def export_job(job_id: str, fmt: str):
    job = JOBS.get(job_id)
    if not job:
        raise HTTPException(404, "unknown job")
    if not job.done.is_set():
        raise HTTPException(409, "job still running")
    if fmt == "zip":
        return _zip_response(job.pins)
    if fmt == "xlsx":
        return _xlsx_response(job.pins)
    raise HTTPException(400, "format must be zip or xlsx")


_SUGGEST_TTL = 300
_suggest_cache: dict[str, tuple[float, list[dict]]] = {}
_suggest_locks: dict[str, threading.Lock] = {}
_suggest_locks_guard = threading.Lock()


@app.get("/api/suggest")
def suggest(q: str = ""):
    q = q.strip().lower()
    if len(q) < 2:
        return {"suggestions": []}
    now = time.monotonic()
    hit = _suggest_cache.get(q)
    if hit and now - hit[0] < _SUGGEST_TTL:
        return {"suggestions": hit[1]}
    with _suggest_locks_guard:
        lock = _suggest_locks.setdefault(q, threading.Lock())
    with lock:
        now = time.monotonic()
        hit = _suggest_cache.get(q)
        if hit and now - hit[0] < _SUGGEST_TTL:
            return {"suggestions": hit[1]}
        try:
            session = build_session()
            out = typeahead_suggestions(session, q)
        except Exception:
            out = []
        _suggest_cache[q] = (now, out)
        if len(_suggest_cache) > 200:
            oldest = sorted(_suggest_cache, key=lambda k: _suggest_cache[k][0])
            for k in oldest[:100]:
                _suggest_cache.pop(k, None)
        return {"suggestions": out}


@app.get("/api/images/{name}")
def global_image(name: str):
    img_dir = (_out_dir() / "images").resolve()
    path = (img_dir / name).resolve()
    if not path.is_relative_to(img_dir) or not path.is_file():
        raise HTTPException(404, "not found")
    return FileResponse(path)


@app.get("/api/gallery")
def get_gallery():
    img_dir = (_out_dir() / "images").resolve()
    if not img_dir.exists():
        return {"pins": [], "total": 0}
    metadata_map: dict[str, dict] = {}
    for jf in _out_dir().glob("*.json"):
        if jf.name == "schedules.json" or jf.name.startswith("."):
            continue
        try:
            items = json.loads(jf.read_text(encoding="utf-8"))
            if isinstance(items, list):
                for p in items:
                    if isinstance(p, dict) and p.get("pin_id"):
                        metadata_map[str(p["pin_id"])] = p
        except Exception:
            pass
    valid_exts = {".jpg", ".jpeg", ".png", ".webp"}
    image_files = sorted(
        [f for f in img_dir.iterdir() if f.is_file() and f.suffix.lower() in valid_exts],
        key=lambda f: f.stat().st_mtime,
        reverse=True,
    )
    gallery_pins = []
    for f in image_files:
        pin_id = f.stem
        meta = metadata_map.get(pin_id)
        if meta:
            p = dict(meta)
            p["local_file"] = f.name
        else:
            p = {
                "pin_id": pin_id,
                "title": f"Pin {pin_id}",
                "description": "",
                "local_file": f.name,
                "image_url": f"/api/images/{f.name}",
                "pin_url": f"https://www.pinterest.com/pin/{pin_id}/" if pin_id.isdigit() else "",
                "saves": None,
                "comments": None,
            }
        gallery_pins.append(p)
    return {"pins": gallery_pins, "total": len(gallery_pins)}


@app.get("/api/gallery/export/zip")
def export_gallery():
    img_dir = (_out_dir() / "images").resolve()
    valid_exts = {".jpg", ".jpeg", ".png", ".webp"}
    pins = (
        [{"local_file": f.name} for f in img_dir.iterdir()
         if f.is_file() and f.suffix.lower() in valid_exts]
        if img_dir.exists() else []
    )
    return _zip_response(pins)


def _zip_response(pins: list[dict]) -> StreamingResponse:
    buf = io.BytesIO()
    img_dir = _out_dir() / "images"
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as zf:
        for p in pins:
            f = p.get("local_file")
            if f and (img_dir / f).is_file():
                zf.write(img_dir / f, f"images/{f}")
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/zip",
                             headers={"Content-Disposition": 'attachment; filename="pins-images.zip"'})


def _xlsx_response(pins: list[dict]) -> StreamingResponse:
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(500, "openpyxl is required for XLSX export")
    wb = Workbook()
    ws = wb.active
    ws.title = "pins"
    cols = list(pins[0].keys()) if pins else ["pin_id"]
    ws.append(cols)
    for p in pins:
        ws.append([str(p.get(c)) if p.get(c) is not None else "" for c in cols])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="pins.xlsx"'},
    )


_visual_cache: dict[str, tuple[float, list[dict]]] = {}


@app.get("/api/visual-search")
def visual_search(pin_id: str = "", limit: int = 25):
    pin_id = pin_id.strip()
    if not pin_id.isdigit():
        raise HTTPException(400, "numeric pin_id required")
    cached = _visual_cache.get(pin_id)
    if cached and time.monotonic() - cached[0] < 600:
        return {"pins": cached[1]}
    try:
        session = build_session()
        pins = related_pins(session, pin_id, limit=min(limit, 50))
    except Exception as e:
        raise HTTPException(502, f"visual search failed: {e}")
    _visual_cache[pin_id] = (time.monotonic(), pins)
    return {"pins": pins}


SCHEDULES_FILE = _out_dir() / "schedules.json"


def _load_schedules() -> list[dict]:
    if SCHEDULES_FILE.exists():
        try:
            return json.loads(SCHEDULES_FILE.read_text(encoding="utf-8"))
        except (ValueError, OSError):
            return []
    return []


def _save_schedules(items: list[dict]) -> None:
    SCHEDULES_FILE.parent.mkdir(parents=True, exist_ok=True)
    SCHEDULES_FILE.write_text(json.dumps(items, ensure_ascii=False, indent=1), encoding="utf-8")


class ScheduleIn(BaseModel):
    mode: str = Field(default="search", pattern="^(search|board)$")
    query: str = Field(min_length=1)
    interval_hours: float = Field(default=24, ge=1, le=720)
    limit: int = Field(default=25, ge=1, le=200)


@app.get("/api/schedules")
def list_schedules():
    return {"schedules": _load_schedules()}


@app.post("/api/schedules")
def add_schedule(sch: ScheduleIn):
    items = _load_schedules()
    entry = {
        "id": uuid.uuid4().hex[:10],
        "mode": sch.mode,
        "query": sch.query.strip(),
        "interval_hours": sch.interval_hours,
        "limit": sch.limit,
        "next_run": time.time() + sch.interval_hours * 3600,
        "created": time.time(),
        "last_run": None,
        "runs": 0,
    }
    items.append(entry)
    _save_schedules(items)
    return entry


@app.delete("/api/schedules/{sid}")
def delete_schedule(sid: str):
    items = [s for s in _load_schedules() if s.get("id") != sid]
    _save_schedules(items)
    return {"ok": True, "remaining": len(items)}


def _scheduler_loop():
    while True:
        try:
            now = time.time()
            changed = False
            schedules = _load_schedules()
            for sch in schedules:
                if now >= sch.get("next_run", now + 3600):
                    req = ScrapeRequest(
                        mode=sch["mode"], query=sch["query"],
                        limit=sch["limit"], download=True, details=True, dedup=True,
                    )
                    job = Job(req, _out_dir())
                    JOBS[job.id] = job
                    threading.Thread(target=job.run, daemon=True).start()
                    sch["last_run"] = now
                    sch["next_run"] = now + sch["interval_hours"] * 3600
                    sch["runs"] = sch.get("runs", 0) + 1
                    sch["last_job_id"] = job.id
                    changed = True
            if changed:
                _save_schedules(schedules)
        except Exception:
            pass
        time.sleep(60)


@app.on_event("startup")
def _start_scheduler():
    threading.Thread(target=_scheduler_loop, daemon=True).start()


class DeleteImagesIn(BaseModel):
    names: list[str] = Field(default_factory=list)
    all: bool = False


@app.post("/api/images/delete")
def delete_images(payload: DeleteImagesIn):
    img_dir = (_out_dir() / "images").resolve()
    if payload.all:
        targets = [p for p in img_dir.glob("*") if p.is_file()]
    else:
        targets = []
        for name in payload.names:
            p = (img_dir / name).resolve()
            if str(p).startswith(str(img_dir)) and p.is_file():
                targets.append(p)
    deleted = 0
    for p in targets:
        try:
            p.unlink()
            deleted += 1
        except OSError:
            pass
    return {"deleted": deleted}


def main():
    import uvicorn
    port = int(os.getenv("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port, log_level="warning")


if __name__ == "__main__":
    main()